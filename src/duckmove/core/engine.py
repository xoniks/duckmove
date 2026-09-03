from __future__ import annotations

import datetime as _dt
import decimal as _decimal
import logging
import os
import re
import tempfile
import threading
from pathlib import Path
from typing import Any, Dict, List, Optional

import duckdb

from .sqlguard import validate_sql

log = logging.getLogger(__name__)

DEFAULT_DB = Path.home() / ".duckmove" / "data.duckdb"
DEFAULT_MAX_ROWS = 500
#: Hard ceiling on rows returned from a single query, regardless of the
#: `max_rows` a caller asks for. Keeps a runaway request from exhausting
#: memory or flooding the model's context.
MAX_ROWS_LIMIT = 10_000
INTERNAL_PREFIX = "_"
SOURCES_TABLE = "_sources"

#: Optional load-path allowlist. When `DUCKMOVE_ALLOWED_DIRS` is set (OS path
#: separator delimited), `load_data` refuses files outside those directories.
#: `run_sql` can never touch the filesystem at all — see `sqlguard`.
ALLOWED_DIRS_ENV = "DUCKMOVE_ALLOWED_DIRS"


class PathNotAllowed(PermissionError):
    """The requested file lies outside the configured allowlist."""


def allowed_dirs() -> List[Path]:
    raw = os.environ.get(ALLOWED_DIRS_ENV, "").strip()
    if not raw:
        return []
    return [Path(p).expanduser().resolve() for p in raw.split(os.pathsep) if p.strip()]


def assert_path_allowed(p: Path) -> None:
    """Raise :class:`PathNotAllowed` if `p` is outside the allowlist.

    A no-op when no allowlist is configured, which is the default: the user
    names the file explicitly, so loading it is the intent. Operators who want
    a hard boundary set `DUCKMOVE_ALLOWED_DIRS`.
    """
    roots = allowed_dirs()
    if not roots:
        return
    target = p.resolve()
    for root in roots:
        try:
            target.relative_to(root)
            return
        except ValueError:
            continue
    raise PathNotAllowed(
        f"'{p}' is outside the allowed directories "
        f"({os.pathsep.join(str(r) for r in roots)})."
    )


READERS = {
    ".csv": "read_csv({p}, header=true)",
    ".tsv": "read_csv({p}, header=true, delim='\t')",
    ".parquet": "read_parquet({p})",
    ".geojson": "ST_Read({p})",
    ".json": "read_json_auto({p})",
    ".shp": "ST_Read({p})",
    ".xlsx": None,  # read_xlsx with openpyxl fallback, see _load_xlsx
}


def geodesic_km_expr(lat_a: str, lon_a: str, lat_b: str, lon_b: str) -> str:
    """SQL expression for WGS84 geodesic distance in km between two points.

    DuckDB spatial's ST_Distance_Spheroid expects POINTs built as
    ST_Point(latitude, longitude) — latitude as X. Verified against
    NYC<->LA ≈ 3,944 km; the (lon, lat) order returns NaN/garbage.
    """
    return (
        f"ST_Distance_Spheroid(ST_Point({lat_a}, {lon_a}), "
        f"ST_Point({lat_b}, {lon_b})) / 1000.0"
    )


def _sanitize_name(name: str) -> str:
    s = re.sub(r"[^0-9a-zA-Z]+", "_", name.strip().lower()).strip("_")
    if not s:
        s = "table"
    if s[0].isdigit():
        s = "t_" + s
    return s


def _jsonify(v: Any) -> Any:
    if isinstance(v, (_dt.date, _dt.datetime, _dt.time)):
        return v.isoformat()
    if isinstance(v, _decimal.Decimal):
        return float(v)
    if isinstance(v, (bytes, bytearray)):
        return v.hex()
    return v


class Engine:
    """A DuckDB connection plus the spatial extension.

    A single DuckDB connection is not safe for concurrent use, so every
    caller must hold :attr:`lock` while touching :attr:`con`. The server
    acquires it once per tool call, in the worker thread.
    """

    def __init__(self, db_path: Optional[str] = None) -> None:
        path = db_path or os.environ.get("DUCKMOVE_DB") or str(DEFAULT_DB)
        if path != ":memory:":
            Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.con = duckdb.connect(path)
        self.con.execute("INSTALL spatial; LOAD spatial;")
        log.info("engine ready (db=%s)", path)

    def close(self) -> None:
        with self.lock:
            self.con.close()

    def __enter__(self) -> Engine:
        return self

    def __exit__(self, *exc: Any) -> None:
        self.close()

    # -- naming -------------------------------------------------------
    def _unique_table_name(self, name: str) -> str:
        base = _sanitize_name(name)
        existing = {t["table"] for t in self.list_tables(include_internal=True)}
        final, i = base, 2
        while final in existing:
            final = f"{base}_{i}"
            i += 1
        return final

    # -- creation -----------------------------------------------------
    def create_table_from_text(self, csv_text: str, name: str) -> Dict[str, Any]:
        table = self._unique_table_name(name)
        with tempfile.NamedTemporaryFile(
            "w", suffix=".csv", delete=False, encoding="utf-8", newline=""
        ) as f:
            f.write(csv_text)
            tmp = f.name
        try:
            self.con.execute(
                f'CREATE TABLE "{table}" AS SELECT * FROM read_csv(?, header=true)',
                [tmp],
            )
        finally:
            os.unlink(tmp)
        return self.describe(table)

    def _create_table(self, p: Path, ext: str, table: str) -> None:
        if ext == ".xlsx":
            self._load_xlsx(p, table)
        else:
            template = READERS[ext]
            if template is None:  # every None-reader needs its own branch above
                raise ValueError(f"No SQL reader is wired up for '{ext}'")
            reader = template.format(p="?")
            self.con.execute(
                f'CREATE TABLE "{table}" AS SELECT * FROM {reader}', [str(p)]
            )

    def create_table_from_file(
        self, path: str, name: Optional[str] = None
    ) -> Dict[str, Any]:
        p = Path(path).expanduser()
        if not p.exists():
            raise FileNotFoundError(f"File not found: {path}")
        assert_path_allowed(p)
        ext = p.suffix.lower()
        if ext not in READERS:
            raise ValueError(
                f"Unsupported extension '{ext}'. Supported: {', '.join(sorted(READERS))}"
            )
        abs_path = str(p.resolve())
        stat = p.stat()
        existing = self._find_source(abs_path)
        if existing and self._table_exists(existing[0]):
            tbl, size, mtime = existing
            if size == stat.st_size and abs(mtime - stat.st_mtime) < 1e-6:
                info = self.describe(tbl)
                info["status"] = "reused"
                return info
            # file changed on disk -> refresh in place, same table name
            self.con.execute(f'DROP TABLE "{tbl}"')
            self._create_table(p, ext, tbl)
            self._record_source(tbl, abs_path, stat)
            info = self.describe(tbl)
            info["status"] = "refreshed"
            return info
        table = self._unique_table_name(name or p.stem)
        self._create_table(p, ext, table)
        self._record_source(table, abs_path, stat)
        info = self.describe(table)
        info["status"] = "created"
        return info

    def rename_table(self, old: str, new: str) -> str:
        self._assert_table(old)
        new_name = _sanitize_name(new)
        if new_name != old and self._table_exists(new_name):
            raise ValueError(f"A table named '{new_name}' already exists")
        self.con.execute(f'ALTER TABLE "{old}" RENAME TO "{new_name}"')
        self._ensure_sources()
        self.con.execute(
            f"UPDATE {SOURCES_TABLE} SET table_name = ? WHERE table_name = ?",
            [new_name, old],
        )
        return new_name

    # -- source tracking (for smart reuse/refresh) -------------------
    def _ensure_sources(self) -> None:
        self.con.execute(
            f"CREATE TABLE IF NOT EXISTS {SOURCES_TABLE} ("
            "table_name VARCHAR PRIMARY KEY, path VARCHAR, size BIGINT, mtime DOUBLE)"
        )

    def _find_source(self, abs_path: str):
        self._ensure_sources()
        return self.con.execute(
            f"SELECT table_name, size, mtime FROM {SOURCES_TABLE} WHERE path = ?",
            [abs_path],
        ).fetchone()

    def _record_source(self, table: str, abs_path: str, stat) -> None:
        self._ensure_sources()
        self.con.execute(
            f"DELETE FROM {SOURCES_TABLE} WHERE table_name = ? OR path = ?",
            [table, abs_path],
        )
        self.con.execute(
            f"INSERT INTO {SOURCES_TABLE} VALUES (?, ?, ?, ?)",
            [table, abs_path, stat.st_size, stat.st_mtime],
        )

    def _table_exists(self, table: str) -> bool:
        return (
            self.con.execute(
                "SELECT 1 FROM information_schema.tables "
                "WHERE table_schema='main' AND table_name = ?",
                [table],
            ).fetchone()
            is not None
        )

    def _load_xlsx(self, p: Path, table: str) -> None:
        try:
            self.con.execute(
                f'CREATE TABLE "{table}" AS SELECT * FROM read_xlsx(?)', [str(p)]
            )
            return
        except duckdb.Error:
            pass  # excel extension unavailable -> openpyxl fallback
        import csv as _csv

        import openpyxl

        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        ws = wb.active
        with tempfile.NamedTemporaryFile(
            "w", suffix=".csv", delete=False, encoding="utf-8", newline=""
        ) as f:
            w = _csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])
            tmp = f.name
        try:
            self.con.execute(
                f'CREATE TABLE "{table}" AS SELECT * FROM read_csv(?, header=true)',
                [tmp],
            )
        finally:
            os.unlink(tmp)

    # -- querying -----------------------------------------------------
    def query(self, sql: str, max_rows: int = DEFAULT_MAX_ROWS) -> Dict[str, Any]:
        body = validate_sql(sql)
        requested = max_rows
        try:
            max_rows = int(max_rows)
        except (TypeError, ValueError):
            max_rows = DEFAULT_MAX_ROWS
        max_rows = max(1, min(max_rows, MAX_ROWS_LIMIT))
        rel = self.con.sql(body)
        cols = list(rel.columns)
        types = [str(t) for t in rel.types]
        select_parts = []
        for c, t in zip(cols, types):
            if t.startswith("GEOMETRY"):
                select_parts.append(f'ST_AsText("{c}") AS "{c}"')
            elif t == "TIMESTAMP WITH TIME ZONE":
                # Fetching TIMESTAMPTZ into Python requires pytz; cast to
                # text so results stay dependency-free and JSON-safe.
                select_parts.append(f'CAST("{c}" AS VARCHAR) AS "{c}"')
            else:
                select_parts.append(f'"{c}"')
        wrapped = rel.project(", ".join(select_parts))
        fetched = wrapped.fetchmany(max_rows + 1)
        truncated = len(fetched) > max_rows
        rows = [{c: _jsonify(v) for c, v in zip(cols, rec)} for rec in fetched[:max_rows]]
        out: Dict[str, Any] = {
            "columns": cols,
            "rows": rows,
            "truncated": truncated,
            "row_count_returned": len(rows),
            "max_rows": max_rows,
        }
        if isinstance(requested, int) and requested > MAX_ROWS_LIMIT:
            out["note"] = (
                f"max_rows was capped at {MAX_ROWS_LIMIT}. Aggregate in SQL "
                f"rather than requesting more rows."
            )
        return out

    # -- inspection ---------------------------------------------------
    def _count(self, table: str) -> int:
        row = self.con.execute(f'SELECT count(*) FROM "{table}"').fetchone()
        return int(row[0]) if row else 0

    def list_tables(self, include_internal: bool = False) -> List[Dict[str, Any]]:
        rows = self.con.execute(
            "SELECT table_name FROM information_schema.tables "
            "WHERE table_schema = 'main' ORDER BY table_name"
        ).fetchall()
        out = []
        for (t,) in rows:
            if not include_internal and t.startswith(INTERNAL_PREFIX):
                continue
            out.append({"table": t, "row_count": self._count(t)})
        return out

    def describe(self, table: str) -> Dict[str, Any]:
        self._assert_table(table)
        cols = self.con.execute(f'DESCRIBE "{table}"').fetchall()
        columns = [{"name": c[0], "type": c[1]} for c in cols]
        return {"table": table, "columns": columns, "row_count": self._count(table)}

    def drop_table(self, table: str) -> None:
        self._assert_table(table)
        self.con.execute(f'DROP TABLE "{table}"')
        self._ensure_sources()
        self.con.execute(f"DELETE FROM {SOURCES_TABLE} WHERE table_name = ?", [table])

    def assert_columns(self, table: str, *cols: str) -> None:
        """Validate identifiers against the actual schema before they are
        interpolated into SQL. Raises KeyError for unknown names."""
        valid = {c["name"] for c in self.describe(table)["columns"]}
        for c in cols:
            if c not in valid:
                raise KeyError(
                    f"Unknown column '{c}' in table '{table}'. "
                    f"Available: {', '.join(sorted(valid))}"
                )

    def _assert_table(self, table: str) -> None:
        ok = self.con.execute(
            "SELECT 1 FROM information_schema.tables "
            "WHERE table_schema='main' AND table_name = ?",
            [table],
        ).fetchone()
        if not ok:
            raise KeyError(f"Unknown table: {table}")
